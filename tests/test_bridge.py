"""Step A: portal submissions feed the OSS report engine unchanged.

Seeds the fictional OSS demo data (slim-report-engine-oss/demo/demo.py) into a temp SQLite file,
converts portal TRs with bridge.to_slim_v1 and runs the report engine's own sheet builder.
Skipped when the OSS packages or the checkout aren't available.
"""

import importlib.util
import sys
from datetime import date, time
from pathlib import Path

import pytest

pytest.importorskip("slim_domain")
pytest.importorskip("slim_report_engine")

import openpyxl  # noqa: E402
from slim_report_engine.cli import build_services  # noqa: E402
from slim_report_engine.reporting import report_writer  # noqa: E402
from slim_report_engine.reporting.submission_report_builder import (  # noqa: E402
    build_submission_sheets,
    sx5_location_from_filename,
)

from slim_lab_portal.bridge import NotReadyForReport, to_slim_v1, v1_file_name  # noqa: E402
from slim_lab_portal.domain import (  # noqa: E402
    PaymentMethod,
    ProcessingTime,
    ReportingUnit,
    RequestType,
    TRStatus,
    WaferSize,
    transition,
)
from slim_lab_portal.oss_guard import assert_oss_database  # noqa: E402
from slim_lab_portal.repositories import TRSubmissionRepository  # noqa: E402
from tests.factories import NOW, STAFF_ACTOR, sample, submission  # noqa: E402

OSS_DEMO = Path(__file__).resolve().parents[2] / "slim-report-engine-oss" / "demo" / "demo.py"


@pytest.fixture(scope="module")
def services(tmp_path_factory):
    if not OSS_DEMO.is_file():
        pytest.skip(f"OSS demo not found at {OSS_DEMO}")
    url = f"sqlite:///{tmp_path_factory.mktemp('oss') / 'demo.db'}"
    assert_oss_database(url)
    spec = importlib.util.spec_from_file_location("slim_oss_demo_for_tests", OSS_DEMO)
    demo = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = demo
    spec.loader.exec_module(demo)
    demo.seed(url)
    return build_services(url)


def _ids(services, *names):
    return tuple(services.analysis_svc.get_by_name(n).id for n in names)


def _received(sub):
    return transition(sub, TRStatus.RECEIVED, actor=STAFF_ACTOR, at=NOW,
                      date_received=date(2026, 10, 9), received_by="JM")


def _sheets(services, v1):
    return build_submission_sheets(
        v1, services.customer_svc.load_customer(v1.customer_id), services.chemical_svc,
        services.analysis_svc, services.element_svc, services.specification_svc,
    )


def _text(sheets) -> str:
    """Writes the sheets with the engine's own report_writer and returns every cell's text."""
    wb = openpyxl.Workbook()
    for sheet in sheets:
        ws = wb.create_sheet(sheet.name)
        report_writer.write_sections(ws, list(sheet.sections), start_row=1, limit_columns=sheet.limit_columns)
    return "\n".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
                     if c.value is not None)


def test_chemical_submission_builds_report_sheets(services):
    chem = services.chemical_svc.get_by_name("Chemical 01")
    sub = _received(submission(
        tr_number="TR00042", id=42, customer_id=2,  # demo customer "B" has specification limits
        samples=(
            sample(1, chemical_id=chem.id, chemical_name=chem.name,
                   analysis_ids=_ids(services, "36 Elements", "pH")),
            sample(2, chemical_id=chem.id, chemical_name=chem.name, analysis_ids=_ids(services, "4 Anions"),
                   processing_time=ProcessingTime.TIME_LIMITED, requested_time=time(15, 0)),
        ),
    ))
    v1 = to_slim_v1(sub)
    assert v1.date_received == date(2026, 10, 9)
    assert v1.technician_initials == "JM"
    assert v1.samples[1].requested_time == "15:00"
    assert v1.file_name.endswith("Chemical TR00042.xlsx")
    content = _text(_sheets(services, v1))
    assert "S-001" in content and "S-002" in content


def test_water_submission_builds_report_sheets(services):
    water = services.chemical_svc.get_by_name("Water")
    sub = _received(submission(
        tr_number="TR00043", id=43, request_type=RequestType.WATER,
        payment_method=PaymentMethod.CREDIT_CARD, po_number="",
        samples=(sample(1, chemical_id=water.id, chemical_name=water.name,
                        analysis_ids=_ids(services, "TOC", "Conductivity")),),
    ))
    v1 = to_slim_v1(sub)
    assert v1.po_information == "" and "lab to call" in v1.credit_card_information
    assert _sheets(services, v1)


@pytest.mark.parametrize("unit", list(ReportingUnit))
def test_wafer_submission_builds_report_sheets(services, unit):
    sub = _received(submission(
        tr_number="TR00044", id=44, request_type=RequestType.WAFER,
        samples=(sample(1, chemical_id=None, chemical_name="", processing_time=ProcessingTime.NEXT_DAY_RUSH,
                        wafer_size=WaferSize.MM_300, reporting_unit=unit,
                        analysis_ids=_ids(services, "36 Elements")),),
    ))
    v1 = to_slim_v1(sub)
    assert v1.samples[0].form_chemical_name == "300mm"
    content = _text(_sheets(services, v1))
    assert "300mm" in content
    # The report engine only recognises its exact unit strings; the scaled one changes the note.
    assert ("Results in units of 1010 atoms" in content) == (unit is ReportingUnit.E10_ATOMS_PER_CM2)


def test_stored_submission_round_trips_into_report(services, db):
    """Through the portal's own storage, as the real flow will be."""
    chem = services.chemical_svc.get_by_name("Chemical 02")
    repo = TRSubmissionRepository(db)
    stored = repo.add(submission(samples=(sample(1, chemical_id=chem.id, chemical_name=chem.name,
                                                 analysis_ids=_ids(services, "36 Elements")),)))
    received = repo.record_transition(stored, _received(stored))
    db.commit()
    assert _sheets(services, to_slim_v1(repo.get_by_tr_number(received.tr_number)))


def test_location_reaches_the_report_engine_via_file_name():
    for location, expected in (("SX5N", "SX5N"), ("SX5S", "SX5S")):
        name = v1_file_name(submission(tr_number="TR00045", location=location))
        assert sx5_location_from_filename(name) == expected


def test_unreceived_submission_is_refused():
    with pytest.raises(NotReadyForReport):
        to_slim_v1(submission(tr_number="TR00046"))


def test_unmatched_chemical_is_refused_until_matched():
    with pytest.raises(NotReadyForReport, match="not yet matched"):
        to_slim_v1(_received(submission(tr_number="TR00047", id=47,
                                        samples=(sample(1, chemical_id=None, chemical_name="IPA"),))))
